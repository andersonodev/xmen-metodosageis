"""Django forms used across the accounts application."""

from __future__ import annotations

import csv
import io
from pathlib import Path

from django import forms
from django.contrib.auth.forms import AuthenticationForm, UserCreationForm
from django.core.exceptions import ValidationError

from .models import User


class UserRegistrationForm(UserCreationForm):
    """Registration form with the additional fields required by the product."""

    first_name = forms.CharField(
        max_length=30,
        required=True,
        widget=forms.TextInput(attrs={
            "class": "form-input",
            "placeholder": "Digite seu primeiro nome",
        }),
        label="Nome",
    )
    last_name = forms.CharField(
        max_length=30,
        required=True,
        widget=forms.TextInput(attrs={
            "class": "form-input",
            "placeholder": "Digite seu sobrenome",
        }),
        label="Sobrenome",
    )
    email = forms.EmailField(
        required=True,
        widget=forms.EmailInput(attrs={
            "class": "form-input",
            "placeholder": "exemplo@empresa.com",
        }),
        label="E-mail",
    )
    username = forms.CharField(
        max_length=150,
        required=True,
        widget=forms.TextInput(attrs={
            "class": "form-input",
            "placeholder": "Nome de usuário único",
        }),
        label="Nome de usuário",
        help_text="Será usado para fazer login no sistema",
    )
    password1 = forms.CharField(
        widget=forms.PasswordInput(attrs={
            "class": "form-input",
            "placeholder": "Digite uma senha segura",
        }),
        label="Senha",
    )
    password2 = forms.CharField(
        widget=forms.PasswordInput(attrs={
            "class": "form-input",
            "placeholder": "Confirme sua senha",
        }),
        label="Confirmar senha",
    )
    role = forms.ChoiceField(
        choices=User.ROLE_CHOICES,
        required=True,
        widget=forms.Select(attrs={"class": "form-select"}),
        label="Função",
        initial="collaborator",
    )
    position = forms.CharField(
        max_length=100,
        required=False,
        widget=forms.TextInput(attrs={
            "class": "form-input",
            "placeholder": "Ex: Desenvolvedor Frontend, Scrum Master...",
        }),
        label="Cargo (opcional)",
    )
    department = forms.CharField(
        max_length=100,
        required=False,
        widget=forms.TextInput(attrs={
            "class": "form-input",
            "placeholder": "Ex: Tecnologia, Produto, Marketing...",
        }),
        label="Departamento (opcional)",
    )

    class Meta:
        model = User
        fields = [
            "username",
            "first_name",
            "last_name",
            "email",
            "password1",
            "password2",
            "role",
            "position",
            "department",
        ]

    def clean_email(self):
        email = self.cleaned_data.get("email")
        if User.objects.filter(email=email).exists():
            raise ValidationError("Este e-mail já está cadastrado.")
        return email

    def clean_username(self):
        username = self.cleaned_data.get("username")
        if User.objects.filter(username=username).exists():
            raise ValidationError("Este nome de usuário já está em uso.")
        return username

    def save(self, commit: bool = True) -> User:
        user = super().save(commit=False)
        user.email = self.cleaned_data["email"]
        user.first_name = self.cleaned_data["first_name"]
        user.last_name = self.cleaned_data["last_name"]
        user.role = self.cleaned_data["role"]
        user.position = self.cleaned_data.get("position", "")
        user.department = self.cleaned_data.get("department", "")

        if commit:
            user.save()
        return user


class UserLoginForm(AuthenticationForm):
    """Custom login form that supports username or e-mail."""

    username = forms.CharField(
        widget=forms.TextInput(
            attrs={
                "class": "form-input",
                "placeholder": "Nome de usuário ou e-mail",
                "autofocus": True,
            }
        ),
        label="Usuário",
        max_length=254,
    )
    password = forms.CharField(
        widget=forms.PasswordInput(attrs={"class": "form-input", "placeholder": "Digite sua senha"}),
        label="Senha",
    )
    remember_me = forms.BooleanField(
        required=False,
        widget=forms.CheckboxInput(attrs={"class": "form-checkbox"}),
        label="Lembrar de mim",
    )

    def clean_username(self):
        username = self.cleaned_data.get("username")
        if username and "@" in username:
            try:
                user = User.objects.get(email=username)
                return user.username
            except User.DoesNotExist:
                pass
        return username

    def clean(self):
        cleaned_data = super().clean()
        user = getattr(self, "user_cache", None)
        if user and not user.is_email_verified:
            raise forms.ValidationError("Confirme seu e-mail antes de continuar.")
        return cleaned_data


class UserProfileForm(forms.ModelForm):
    """Form used for collaborator profile updates."""

    first_name = forms.CharField(
        max_length=30,
        widget=forms.TextInput(attrs={"class": "form-input"}),
        label="Nome",
    )
    last_name = forms.CharField(
        max_length=30,
        widget=forms.TextInput(attrs={"class": "form-input"}),
        label="Sobrenome",
    )
    email = forms.EmailField(
        widget=forms.EmailInput(attrs={"class": "form-input"}),
        label="E-mail",
    )
    bio = forms.CharField(
        required=False,
        widget=forms.Textarea(attrs={"class": "form-textarea", "rows": 4}),
        label="Biografia",
    )
    position = forms.CharField(
        required=False,
        widget=forms.TextInput(attrs={"class": "form-input"}),
        label="Cargo",
    )
    department = forms.CharField(
        required=False,
        widget=forms.TextInput(attrs={"class": "form-input"}),
        label="Departamento",
    )

    class Meta:
        model = User
        fields = [
            "first_name",
            "last_name",
            "email",
            "bio",
            "avatar",
            "position",
            "department",
            "availability_hours",
            "is_available",
        ]

    def save(self, commit: bool = True) -> User:
        user = super().save(commit=False)
        if commit:
            user.save()
        return user


class CollaboratorImportForm(forms.Form):
    file = forms.FileField(
        label="Arquivo CSV ou Excel",
        widget=forms.ClearableFileInput(attrs={"class": "form-input"}),
    )
    default_role = forms.ChoiceField(
        choices=User.ROLE_CHOICES,
        initial="collaborator",
        label="Papel padrão",
        widget=forms.Select(attrs={"class": "form-select"}),
    )
    send_welcome_email = forms.BooleanField(
        required=False,
        initial=True,
        label="Simular envio de e-mail de boas-vindas",
        widget=forms.CheckboxInput(attrs={"class": "form-checkbox"}),
    )

    def clean_file(self):
        file = self.cleaned_data["file"]
        extension = Path(file.name).suffix.lower()
        if extension not in {".csv", ".xlsx"}:
            raise ValidationError("Envie um arquivo CSV ou XLSX válido.")
        return file

    def iter_rows(self):
        file = self.cleaned_data["file"]
        extension = Path(file.name).suffix.lower()
        file.seek(0)

        if extension == ".csv":
            text_file = io.TextIOWrapper(file, encoding="utf-8-sig")
            reader = csv.DictReader(text_file)
            for row in reader:
                yield row
        else:
            try:
                from openpyxl import load_workbook
            except ImportError as exc:  # pragma: no cover - dependency missing
                raise ValidationError("Biblioteca openpyxl não disponível.") from exc

            workbook = load_workbook(file, read_only=True)
            sheet = workbook.active
            headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if any(row):
                    yield {header: value for header, value in zip(headers, row)}


class UserRoleForm(forms.ModelForm):
    class Meta:
        model = User
        fields = ["role", "is_active"]
        widgets = {
            "role": forms.Select(attrs={"class": "form-select"}),
            "is_active": forms.CheckboxInput(attrs={"class": "form-checkbox"}),
        }
